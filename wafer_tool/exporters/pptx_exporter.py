"""
exporters/pptx_exporter.py
==========================
Generates a PowerPoint report from ordered PNG snapshots of the PDF pages.

Gracefully degrades: if python-pptx is not installed, HAS_PPTX is False
and create_powerpoint_report raises ImportError with install instructions.

Template resolution order
--------------------------
1. <PyInstaller _MEIPASS>/Infineon Default Theme.pptx
2. <executable directory>/Infineon Default Theme.pptx
3. <this file's directory>/Infineon Default Theme.pptx
4. python-pptx blank presentation (fallback)

Public API
----------
HAS_PPTX : bool
create_powerpoint_report(page_png_dir, pptx_path, title_text)
"""
from __future__ import annotations

import os
import re
import sys

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.util import Pt
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

__all__ = ["HAS_PPTX", "create_powerpoint_report"]

TEMPLATE_FILENAME = "Infineon Default Theme.pptx"


def create_powerpoint_report(
    page_png_dir: str,
    pptx_path: str,
    title_text: str,
) -> None:
    """
    Compile ordered PNG snapshots of PDF-style report pages into a .pptx file.

    Parameters
    ----------
    page_png_dir : Directory containing ordered report-page PNG files.
    pptx_path    : Destination path for the generated PowerPoint file.
    title_text   : Text displayed on the title slide.

    Raises
    ------
    ImportError
        If python-pptx is not installed.
    """
    if not HAS_PPTX:
        raise ImportError(
            "python-pptx is not installed.\n"
            "Install it with:  pip install python-pptx"
        )

    prs = _load_presentation()
    content_layout = _content_slide_layout(prs)

    # ── Title slide ───────────────────────────────────────────────────────
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    if title_slide.shapes.title:
        title_slide.shapes.title.text = title_text

    # ── Report page slides ────────────────────────────────────────────────
    # Sort by the LEADING PAGE NUMBER numerically, not lexicographically: the
    # page-image names are zero-padded to only 3 digits, so a plain string sort
    # puts "1000_..." before "100_..." and "999_..." — scrambling every slide
    # past page 999 (e.g. a 2000+ page report from a large multi-lot file).
    def _page_order(fname: str):
        m = re.match(r"(\d+)", fname)
        return (int(m.group(1)) if m else 1 << 30, fname)

    png_files = sorted(
        (f for f in os.listdir(page_png_dir) if f.lower().endswith(".png")),
        key=_page_order,
    )
    for filename in png_files:
        img_path = os.path.join(page_png_dir, filename)
        slide = prs.slides.add_slide(content_layout)
        _remove_placeholders(slide, _outline_title(filename))
        _add_report_page(slide, prs, img_path)

    import gc; gc.collect()   # free any lingering PNG buffers before writing
    prs.save(pptx_path)
    print(f"✓ PowerPoint saved: {pptx_path}")


# ---------------------------------------------------------------------------
# Private helpers
# ---------------------------------------------------------------------------

def _find_template() -> str | None:
    """Return the first existing template path, or None."""
    candidates: list[str] = []
    if getattr(sys, "frozen", False):
        candidates.append(os.path.join(sys._MEIPASS, TEMPLATE_FILENAME))
        candidates.append(os.path.join(os.path.dirname(sys.executable), TEMPLATE_FILENAME))
    candidates.append(
        os.path.join(os.path.dirname(os.path.abspath(__file__)), TEMPLATE_FILENAME)
    )
    return next((p for p in candidates if os.path.exists(p)), None)


def _load_presentation() -> "Presentation":
    template = _find_template()
    if template:
        try:
            prs = Presentation(template)
            print(f"✓ Loaded PPTX template: {template}")
            _ensure_widescreen(prs)
            return prs
        except Exception as exc:
            print(f"⚠  Template load failed ({exc}). Using default.")
    else:
        print("⚠  No PPTX template found — using python-pptx default.")
    prs = Presentation()
    _ensure_widescreen(prs)
    return prs


def _ensure_widescreen(prs: "Presentation") -> None:
    """Force a 16:9 widescreen slide (13.333 × 7.5 in).

    python-pptx's default template is 4:3 (10 × 7.5 in); on a widescreen monitor
    PowerPoint letterboxes it with grey bars, making every slide (and its plot)
    look small. Widescreen fills modern displays. Only resizes when the deck is
    not already ~16:9 so a genuinely 16:9 theme template is left untouched.
    """
    target_w, target_h = 12192000, 6858000   # 13.333 × 7.5 in in EMU
    try:
        if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.02:
            prs.slide_width = target_w
            prs.slide_height = target_h
    except Exception:
        pass



def _content_slide_layout(prs: "Presentation"):
    """Return the layout used for report-page slides (placeholders are stripped)."""
    try:
        return prs.slide_layouts[5]
    except IndexError:
        return prs.slide_layouts[-1]


def _add_report_page(slide, prs: "Presentation", img_path: str) -> None:
    """Place a PDF-style page snapshot scaled to fill the slide (fit-to-box,
    aspect preserved) and centred."""
    margin = Pt(10)
    top = margin
    avail_w = prs.slide_width - 2 * margin
    avail_h = prs.slide_height - top - margin

    # Add at native size, then scale by whichever dimension is the tighter fit
    # so the image is as large as possible without overflowing or distorting.
    picture = slide.shapes.add_picture(img_path, 0, 0)
    if picture.width and picture.height:
        scale = min(avail_w / picture.width, avail_h / picture.height)
        picture.width = int(picture.width * scale)
        picture.height = int(picture.height * scale)
    picture.left = int((prs.slide_width - picture.width) / 2)
    picture.top = int(top + (avail_h - picture.height) / 2)


def _outline_title(filename: str) -> str:
    """Return a readable title for PowerPoint's Outline pane."""
    stem = os.path.splitext(filename)[0]
    stem = re.sub(r"^\d+_", "", stem)
    lot_match = re.fullmatch(r"lot_(.+)_page_(\d+)", stem)
    if lot_match:
        return f"Lot ID: {lot_match.group(1)} | Page {lot_match.group(2)}"
    stem = re.sub(r"_page_(\d+)$", r" - Page \1", stem)
    return stem.replace("_", " ").title()


def _remove_placeholders(slide, title_text: str) -> None:
    # Retain an off-slide title placeholder so PowerPoint shows page names in Outline.
    for ph in list(slide.placeholders):
        if ph.placeholder_format.type in {
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        }:
            ph.text = title_text
            ph.left = -abs(ph.width)
            ph.top = -abs(ph.height)
            for paragraph in ph.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(1)
                    run.font.color.rgb = RGBColor(255, 255, 255)
        else:
            ph._element.getparent().remove(ph._element)
